# Copyright 2026 会议库 contributors
# SPDX-License-Identifier: Apache-2.0

"""Word（.docx）渲染：python-docx，出处作小字脚注式补充。"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from huiyiku.export.common import ms_to_clock


def _new_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    # 中文字形必须单独设 eastAsia，否则 Word 里可能回退成默认字体
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    # 标题样式有独立的字体设置，不继承 Normal，也要逐个设一遍
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return doc


def _small_gray(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x76, 0x7E, 0x8D)


def _blue(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x4F, 0x6E, 0xF7)


def _add_item(doc: Document, item: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    p.add_run(item.get("text") or "")
    if item.get("status") == "unsupported":
        run = p.add_run("  〔AI 补充·录音未找到〕")
        run.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
    elif item.get("status") == "quoted":
        run = p.add_run("  〔录音原文〕")
        run.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
    if item.get("owner_name"):
        _small_gray(
            p,
            f"　负责人：{item['owner_name']}"
            + (f"，截止：{item['due_at']}" if item.get("due_at") else ""),
        )
    ev = item.get("evidence") or []
    if ev:
        ev_p = doc.add_paragraph()
        _blue(
            ev_p,
            "出处："
            + "；".join(
                f"{ms_to_clock(e.get('start_ms'))}"
                + ("（时间近似）" if e.get("vector_matched") else "")
                + f" · {e.get('speaker_name') or '?'}"
                for e in ev
            ),
        )
        for e in ev:
            q = doc.add_paragraph()
            _small_gray(
                q,
                f"{ms_to_clock(e.get('start_ms'))} {e.get('speaker_name') or '?'}：{e.get('text') or ''}",
            )
    else:
        note = doc.add_paragraph()
        warn = note.add_run("出处：未在录音中定位到原文（AI 补充）")
        warn.font.size = Pt(8.5)
        warn.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)


def _add_meeting(doc: Document, mdoc: dict[str, Any]) -> None:
    doc.add_heading(mdoc["title"], level=1)
    status = "已接受" if mdoc.get("version_status") == "accepted" else "草稿"
    meta = doc.add_paragraph()
    _small_gray(meta, f"时间：{mdoc.get('occurred_at') or ''}　报告状态：{status}")
    if mdoc.get("stale_reason"):
        note = doc.add_paragraph()
        run = note.add_run(mdoc["stale_reason"])
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if mdoc.get("auto_speakers_unreviewed"):
        note = doc.add_paragraph()
        _small_gray(note, "说话人尚未人工校对，出处里的人物名为机器标注。")
    if mdoc.get("summary"):
        doc.add_heading("摘要", level=2)
        doc.add_paragraph(mdoc["summary"])
    for sec in mdoc.get("sections") or []:
        doc.add_heading(sec["title"], level=2)
        for item in sec["items"]:
            _add_item(doc, item)
    if mdoc.get("participants"):
        doc.add_heading("参会人", level=2)
        names = "、".join(
            str(p.get("name") or "?")
            + (f"（发言 {p['segment_count']} 段）" if p.get("segment_count") else "")
            for p in mdoc["participants"]
            if isinstance(p, dict)
        )
        doc.add_paragraph(names)
    transcript = mdoc.get("transcript")
    if transcript:
        doc.add_heading("转写", level=2)
        for t in transcript:
            p = doc.add_paragraph()
            _blue(p, f"{ms_to_clock(t['start_ms'])} ")
            p.add_run(f"{t['speaker_name']}：{t['text']}")
    if not mdoc.get("has_report"):
        _small_gray(doc.add_paragraph(), "（本会议还没有生成报告）")


def render_meeting_docx(mdoc: dict[str, Any]) -> bytes:
    doc = _new_document()
    doc.add_heading(mdoc["title"], level=0)
    _add_meeting(doc, mdoc)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def render_group_docx(
    group_name: str, digest_md: str, meeting_docs: list[dict[str, Any]]
) -> bytes:
    doc = _new_document()
    doc.add_heading(f"会议组汇总：{group_name}", level=0)
    _small_gray(doc.add_paragraph(), f"共 {len(meeting_docs)} 场会议")
    if digest_md:
        doc.add_heading("项目进展", level=1)
        for line in digest_md.splitlines():
            if line.strip():
                doc.add_paragraph(line)
    for mdoc in meeting_docs:
        _add_meeting(doc, mdoc)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
