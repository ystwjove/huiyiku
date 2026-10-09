# Copyright 2026 会议库 contributors
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

import csv
import io
import json

from fastapi.testclient import TestClient
from sqlalchemy import select

from huiyiku.api.app import create_app
from huiyiku.db.models import Report, ReportVersion
from huiyiku.db.session import make_engine, session_factory
from huiyiku.store import create_meeting, persist_transcript
from huiyiku.timeutil import now_iso


def _seed(runtime, *, with_report: bool) -> int:
    engine = make_engine(runtime.layout["db"])
    Session = session_factory(engine)
    with Session() as session:
        m = create_meeting(session, title="导出会")
        persist_transcript(
            session,
            m,
            [
                {
                    "start_ms": 0,
                    "end_ms": 60000,
                    "text": "项目组决定在下个季度完成预算审批并启动供应商招标流程",
                    "speaker_label": "SPEAKER_00",
                },
                {
                    "start_ms": 60000,
                    "end_ms": 120000,
                    "text": "另外团队计划去滨海城市组织团建活动",
                    "speaker_label": "SPEAKER_01",
                },
            ],
            provider="funasr",
            model="paraformer-zh",
            version="v2.0.4",
        )
        if with_report:
            ts = now_iso()
            report = Report(meeting_id=m.id, created_at=ts, updated_at=ts)
            session.add(report)
            session.flush()
            data = {
                "summary": "一次关于预算与团建的讨论",
                "key_points": [],
                "decisions": [
                    {"content": "完成预算审批并启动供应商招标流程", "segment_ids": []},
                    {"content": "团队决定去火星建立殖民地", "segment_ids": []},
                ],
                "action_items": [
                    {
                        "title": "完成预算审批并启动供应商招标流程",
                        "owner_person_id": None,
                        "due_at": None,
                        "priority": "high",
                        "segment_ids": [],
                    }
                ],
                "risks": [],
                "open_questions": [],
                "participants": [],
            }
            ver = ReportVersion(
                report_id=report.id,
                version=1,
                source="llm",
                content_json=json.dumps(data, ensure_ascii=False),
                status="draft",
                created_at=ts,
            )
            session.add(ver)
            session.flush()
            report.current_version_id = ver.id
        session.commit()
        return m.id


def test_export_md_contains_provenance(runtime) -> None:
    mid = _seed(runtime, with_report=True)
    client = TestClient(create_app(runtime))
    r = client.get(f"/api/export/meeting/{mid}.md")
    assert r.status_code == 200
    text = r.content.decode("utf-8")
    assert "导出会" in text
    assert "预算审批" in text
    assert "出处：00:00:00" in text
    assert "未在录音中定位到原文（AI 补充）" in text
    assert "attachment" in r.headers.get("content-disposition", "")
    assert "转写" in text


def test_export_md_without_report_still_works(runtime) -> None:
    mid = _seed(runtime, with_report=False)
    client = TestClient(create_app(runtime))
    r = client.get(f"/api/export/meeting/{mid}.md")
    assert r.status_code == 200
    assert "还没有生成报告" in r.content.decode("utf-8")


def test_export_html_is_self_contained(runtime) -> None:
    mid = _seed(runtime, with_report=True)
    client = TestClient(create_app(runtime))
    r = client.get(f"/api/export/meeting/{mid}.html")
    assert r.status_code == 200
    html = r.content.decode("utf-8")
    assert html.startswith("<!doctype html>")
    assert "<style>" in html and "</style>" in html
    assert "出处" in html and "AI 补充" in html
    assert "script" not in html  # 无 JS，打印友好


def test_export_docx_roundtrip(runtime) -> None:
    mid = _seed(runtime, with_report=True)
    client = TestClient(create_app(runtime))
    r = client.get(f"/api/export/meeting/{mid}.docx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    import docx

    d = docx.Document(io.BytesIO(r.content))
    joined = "\n".join(p.text for p in d.paragraphs)
    assert "导出会" in joined
    assert "预算审批" in joined
    assert "出处：" in joined


def test_export_actions_csv_has_bom_and_provenance(runtime) -> None:
    mid = _seed(runtime, with_report=True)
    engine = make_engine(runtime.layout["db"])
    Session = session_factory(engine)
    with Session() as session:
        from huiyiku.db.models import Meeting
        from huiyiku.store import accept_report_actions

        meeting = session.get(Meeting, mid)
        ver = session.scalar(
            select(ReportVersion).join(Report).where(Report.meeting_id == mid)
        )
        accept_report_actions(session, meeting, ver)
        session.commit()

    client = TestClient(create_app(runtime))
    r = client.get("/api/export/actions.csv")
    assert r.status_code == 200
    raw = r.content
    assert raw.startswith(b"\xef\xbb\xbf"), "CSV 必须带 UTF-8 BOM（Excel 兼容）"
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert len(rows) == 1
    row = rows[0]
    assert row["行动项"] == "完成预算审批并启动供应商招标流程"
    assert row["出处时间点"].startswith("00:00:00")
    assert row["出处人物"]


def test_export_group_markdown(runtime) -> None:
    _seed(runtime, with_report=True)
    client = TestClient(create_app(runtime))
    groups = client.get("/api/groups").json()
    gid = next(g["id"] for g in groups)
    r = client.get(f"/api/export/group/{gid}.md")
    assert r.status_code == 200
    text = r.content.decode("utf-8")
    assert "会议组汇总" in text
    assert "出处" in text
    # 汇总不含转写全文
    assert "## 转写" not in text
    bad = client.get("/api/export/group/999.md")
    assert bad.status_code == 404


def test_export_filename_blocks_windows_illegal_chars() -> None:
    from huiyiku.export.common import export_filename

    # * ? " < > | : 及控制字符都替换，不再 500 或落成 NTFS 备用数据流
    name = export_filename('排期*会"第:二期?', "2026-09-28T10:00:00+08:00", "docx")
    illegal = ["*", "?", chr(34), "<", ">", "|", ":", chr(92)]
    for ch in illegal:
        assert ch not in name
    assert name.endswith("-20260928.docx")
    # 标题是 "." 时 safe_filename 会抛错，导出要兜底
    assert export_filename(".", "2026-09-01", "md").endswith("-20260901.md")


def test_actions_csv_formula_injection_guard() -> None:
    from huiyiku.export.csvout import render_actions_csv

    text = render_actions_csv([{"行动项": "=HYPERLINK(\"http://x\")", "负责人": "+123"}])
    assert "'=" in text and "'+" in text
    assert "=HYPERLINK" not in text.replace("'=", "")


def test_export_filename_reserved_device_names() -> None:
    """Windows 保留设备名不能作为文件名主体（实测 aux 会 OSError）。"""
    from huiyiku.export.common import export_filename

    for title in ("aux", "CON", "nul", "COM1", "lpt9", "aux.txt"):
        name = export_filename(title, "2026-09-28", "md")
        stem = name.split(".")[0].lower()
        assert stem not in {
            "con", "prn", "aux", "nul",
            *(f"com{i}" for i in range(1, 10)),
            *(f"lpt{i}" for i in range(1, 10)),
        }, name
        assert name.endswith("-20260928.md")


def test_ascii_fallback_name_is_readable() -> None:
    """纯中文标题的 filename= 回退名不能退化成 "-20260924.md"。"""
    from huiyiku.export.common import ascii_fallback_name

    assert ascii_fallback_name("海信数据沟通-20260924.md") == "meeting-20260924.md"
    assert ascii_fallback_name("Q3-复盘-20260924.docx") == "Q3-20260924.docx"
    assert ascii_fallback_name("报告") == "meeting"


def test_content_disposition_has_readable_ascii_fallback(runtime) -> None:
    mid = _seed(runtime, with_report=True)
    client = TestClient(create_app(runtime))
    r = client.get(f"/api/export/meeting/{mid}.md")
    cd = r.headers["content-disposition"]
    assert 'filename="' in cd and "meeting-" in cd
    assert "filename*=UTF-8''" in cd


def test_group_html_and_docx_render(runtime) -> None:
    _seed(runtime, with_report=True)
    client = TestClient(create_app(runtime))
    gid = next(g["id"] for g in client.get("/api/groups").json())
    html = client.get(f"/api/export/group/{gid}.html")
    assert html.status_code == 200
    body = html.content.decode("utf-8")
    assert "会议组汇总" in body and "<style>" in body and "<script" not in body
    docx_res = client.get(f"/api/export/group/{gid}.docx")
    assert docx_res.status_code == 200
    import docx

    joined = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_res.content)).paragraphs)
    assert "会议组汇总" in joined


def test_actions_csv_date_filters(runtime) -> None:
    from huiyiku.db.models import Meeting as _M
    from huiyiku.db.session import make_engine, session_factory
    from huiyiku.export.csvout import build_action_rows, render_actions_csv
    from huiyiku.store import accept_report_actions

    mid = _seed(runtime, with_report=True)
    engine = make_engine(runtime.layout["db"])
    Session = session_factory(engine)
    with Session() as session:
        meeting = session.get(_M, mid)
        ver = session.scalar(select(ReportVersion).join(Report).where(Report.meeting_id == mid))
        accept_report_actions(session, meeting, ver)
        session.commit()
        day = (meeting.occurred_at or "")[:10]
        assert build_action_rows(session, date_from=day, date_to=day)
        assert build_action_rows(session, date_from="2999-01-01") == []
        assert build_action_rows(session, date_to="1970-01-01") == []
    text = render_actions_csv([])
    assert text.startswith("会议,")  # 表头仍在，空结果不报错


def test_export_qa_thread(runtime) -> None:
    """问询记录导出：Q/A、AI 补充拆块、引用、附件头、空线程 404。"""
    from huiyiku.db.models import QaMessage
    from huiyiku.db.session import make_engine, session_factory
    from huiyiku.llm.qa import EXTRA_MARKER

    engine = make_engine(runtime.layout["db"])
    Session = session_factory(engine)
    with Session() as session:
        session.add(
            QaMessage(
                session_id="thread-export",
                scope_json="{}",
                question="预算审批定在什么时候？",
                answer="定在下个季度。\n\n" + EXTRA_MARKER + "\n背景：预算流程常见于财年。",
                citations_json=(
                    '[{"start_ms": 120000, "speaker_name": "张三", "text": "定在下个季度",'
                    ' "chunk_type": "transcript"},'
                    ' {"start_ms": 0, "text": "摘要条目", "chunk_type": "summary"}]'
                ),
                retrieval_json='{"allow_extra": true}',
                model_json="{}",
                created_at="2026-09-28T10:00:00+08:00",
            )
        )
        session.commit()

    client = TestClient(create_app(runtime))
    r = client.get("/api/export/qa/thread-export.md")
    assert r.status_code == 200
    text = r.content.decode("utf-8")
    assert "预算审批定在什么时候" in text
    assert "定在下个季度" in text
    assert "AI 补充（非录音内容" in text and "背景：预算流程" in text
    assert "00:02:00 · 张三" in text
    assert "报告条目：摘要条目" in text
    assert "attachment" in r.headers.get("content-disposition", "")

    h = client.get("/api/export/qa/thread-export.html")
    assert h.status_code == 200
    body = h.content.decode("utf-8")
    assert "<style>" in body and "背景：预算流程" in body

    assert client.get("/api/export/qa/thread-none.md").status_code == 404


def test_export_qa_html_escapes_user_content(runtime) -> None:
    """自包含 HTML 导出本地打开：问题/答案/引用里的 HTML 必须转义。"""
    from fastapi.testclient import TestClient

    from huiyiku.api.app import create_app
    from huiyiku.db.models import QaMessage
    from huiyiku.db.session import make_engine, session_factory

    engine = make_engine(runtime.layout["db"])
    Session = session_factory(engine)
    with Session() as session:
        session.add(
            QaMessage(
                session_id="thread-xss",
                scope_json="{}",
                question='<script>alert("q")</script>',
                answer='<img src=x onerror=alert(1)>',
                citations_json=(
                    '[{"start_ms": 1000, "speaker_name": "<b>张三</b>", "text": "<i>原文</i>",'
                    ' "chunk_type": "transcript"},'
                    ' {"start_ms": 0, "text": "<u>报告条目</u>", "chunk_type": "summary"}]'
                ),
                retrieval_json="{}",
                model_json="{}",
                created_at="2026-09-28T10:00:00+08:00",
            )
        )
        session.commit()

    client = TestClient(create_app(runtime))
    body = client.get("/api/export/qa/thread-xss.html").content.decode("utf-8")
    for raw in ("<script>", "<img", "<b>张三</b>", "<i>原文</i>", "<u>报告条目</u>"):
        assert raw not in body, f"未转义: {raw}"
    assert "&lt;script&gt;" in body
    assert "&lt;b&gt;张三&lt;/b&gt;" in body
